# -*- coding: utf-8 -*-
"""Genera el instructivo compartible para construir Portal RISKO en macOS."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
PORTAL_DIR = BASE_DIR.parent
OUTPUT_PATH = BASE_DIR / "Instructivo paso a paso - Construccion Instalador RISKO Mac.docx"
ICON_PATH = PORTAL_DIR / "assets" / "portal-risko-256.png"

PRIMARY = "003C86"
DEEP = "081F3D"
ACCENT = "F7BB26"
LIGHT_BLUE = "EAF2FB"
LIGHT_GOLD = "FFF5D8"
LIGHT_GRAY = "F4F7FB"
LINE = "D5DFEC"
TEXT = "111C2D"
MUTED = "506078"
DANGER = "B42318"
SUCCESS = "0F8A7A"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color: str = LINE, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), PRIMARY)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(color)
    properties.append(underline)
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 28, PRIMARY),
        ("Subtitle", 14, MUTED),
        ("Heading 1", 18, PRIMARY),
        ("Heading 2", 13, DEEP),
        ("Heading 3", 11, PRIMARY),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("BANCO DE BOGOTÁ  |  RIESGO DE MERCADO  |  PORTAL RISKO")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(PRIMARY)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = footer_p.add_run("Uso interno · Instructivo macOS ARM64 · Versión 1.0")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(footer.add_paragraph())


def add_cover(document: Document) -> None:
    if ICON_PATH.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run().add_picture(str(ICON_PATH), width=Cm(2.4))

    accent = document.add_table(rows=1, cols=1)
    accent.autofit = False
    accent.columns[0].width = Cm(16.8)
    cell = accent.cell(0, 0)
    set_cell_shading(cell, ACCENT)
    set_cell_margins(cell, top=70, bottom=70)
    cell.paragraphs[0].add_run(" ")

    document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.add_run("Instructivo paso a paso")
    subtitle = document.add_paragraph(style="Title")
    subtitle.add_run("Construcción e instalación de Portal RISKO en Mac")

    badge = document.add_table(rows=1, cols=1)
    badge.alignment = WD_TABLE_ALIGNMENT.LEFT
    badge_cell = badge.cell(0, 0)
    set_cell_shading(badge_cell, LIGHT_BLUE)
    set_cell_border(badge_cell, PRIMARY, "6")
    set_cell_margins(badge_cell, top=110, bottom=110, start=180, end=180)
    p = badge_cell.paragraphs[0]
    run = p.add_run("macOS ARM64  |  Apple Silicon M1, M2, M3 o posterior")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(PRIMARY)

    document.add_paragraph()
    p = document.add_paragraph()
    p.add_run("Objetivo: ").bold = True
    p.add_run(
        "guiar a una persona sin experiencia previa en macOS para construir, "
        "publicar e instalar la versión piloto de Portal RISKO para Mac."
    )

    document.add_paragraph()
    control = document.add_table(rows=5, cols=2)
    control.alignment = WD_TABLE_ALIGNMENT.LEFT
    control_data = (
        ("Documento", "Instructivo de construcción del instalador Mac"),
        ("Versión", "1.0"),
        ("Fecha", "20 de agosto de 2026"),
        ("Alcance", "Piloto interno · macOS ARM64"),
        ("Resultado", "Instalador RISKO Mac.pkg"),
    )
    for row, (label, value) in zip(control.rows, control_data):
        prevent_row_split(row)
        left, right = row.cells
        set_cell_shading(left, PRIMARY)
        set_cell_shading(right, LIGHT_GRAY)
        set_cell_border(left)
        set_cell_border(right)
        set_cell_margins(left)
        set_cell_margins(right)
        label_run = left.paragraphs[0].add_run(label)
        label_run.bold = True
        label_run.font.color.rgb = RGBColor(255, 255, 255)
        right.paragraphs[0].add_run(value)

    document.add_paragraph()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("USO INTERNO")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(DANGER)
    document.add_page_break()


def add_callout(document: Document, title: str, text: str, *, kind: str = "info") -> None:
    settings = {
        "info": (LIGHT_BLUE, PRIMARY, "INFORMACIÓN"),
        "warning": (LIGHT_GOLD, "9A6700", "IMPORTANTE"),
        "danger": ("FDECEC", DANGER, "DETÉNGASE"),
        "success": ("E8F6F3", SUCCESS, "RESULTADO ESPERADO"),
    }
    fill, color, default_label = settings[kind]
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color, "10")
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    p = cell.paragraphs[0]
    label = p.add_run(f"{title or default_label}: ")
    label.bold = True
    label.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_command(document: Document, command: str, expected: str | None = None) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, DEEP)
    set_cell_border(cell, DEEP, "4")
    set_cell_margins(cell, top=120, bottom=120, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(command.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    if expected:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.3)
        label = p.add_run("Debe mostrar: ")
        label.bold = True
        label.font.color.rgb = RGBColor.from_string(SUCCESS)
        result = p.add_run(expected)
        result.font.name = "Consolas"
        result.font.size = Pt(9.5)


def add_bullets(document: Document, items: list[str], *, checkbox: bool = False) -> None:
    for item in items:
        if checkbox:
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.add_run(f"☐  {item}")
        else:
            document.add_paragraph(item, style="List Bullet")


def add_numbered_actions(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_step_heading(document: Document, number: int, title: str) -> None:
    heading = document.add_heading(level=1)
    run = heading.add_run(f"Paso {number}. {title}")
    run.font.color.rgb = RGBColor.from_string(PRIMARY)


def build_document() -> Document:
    document = Document()
    style_document(document)
    add_cover(document)

    document.add_heading("Cómo utilizar este instructivo", level=1)
    document.add_paragraph(
        "Siga los pasos en orden. Todos los comandos se pueden copiar y pegar; "
        "no es necesario escribirlos manualmente. En macOS, la tecla ⌘ se llama "
        "Command y la tecla Enter también puede aparecer como Return."
    )
    add_callout(
        document,
        "Antes de comenzar",
        "Conecte el Mac a la corriente y no cierre Terminal mientras se esté construyendo el instalador.",
        kind="warning",
    )

    document.add_heading("Lista de preparación", level=2)
    add_bullets(
        document,
        [
            "Mac con chip Apple M1, M2, M3 o posterior.",
            "Conexión a la red de la oficina o a la VPN corporativa.",
            "Usuario y contraseña corporativos para acceder a las carpetas compartidas.",
            "Contraseña de administrador del Mac o acompañamiento de TI.",
            "Acceso a internet para instalar Python y descargar dependencias.",
            "Permiso de escritura en Portal Riesgos de Mercado para publicar el resultado.",
        ],
        checkbox=True,
    )
    add_callout(
        document,
        "Permisos",
        "Si el Mac no permite instalar programas o solicita credenciales que no posee, deténgase y solicite acompañamiento de TI. No comparta su contraseña.",
        kind="danger",
    )

    document.add_heading("Resumen del proceso", level=2)
    summary = document.add_table(rows=1, cols=3)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.autofit = True
    headers = ("Fase", "Qué se hace", "Resultado")
    for cell, value in zip(summary.rows[0].cells, headers):
        set_cell_shading(cell, PRIMARY)
        set_cell_border(cell)
        set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(summary.rows[0])
    for values in (
        ("1. Preparar", "Conectar red e instalar herramientas", "Mac listo para construir"),
        ("2. Construir", "Ejecutar el constructor ARM64", "Instalador RISKO Mac.pkg"),
        ("3. Publicar", "Copiar release y manifiesto Mac", "Actualización automática habilitada"),
        ("4. Probar", "Instalar y abrir Portal RISKO", "Piloto operativo"),
    ):
        row = summary.add_row()
        prevent_row_split(row)
        for cell, value in zip(row.cells, values):
            set_cell_shading(cell, "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].add_run(value)

    add_step_heading(document, 1, "Confirmar que el Mac es ARM")
    add_numbered_actions(
        document,
        [
            "Pulse el símbolo de Apple  situado arriba a la izquierda.",
            "Seleccione Acerca de este Mac.",
            "Busque el campo Chip.",
            "Confirme que dice Apple M1, Apple M2, Apple M3 o una generación posterior.",
        ],
    )
    add_callout(
        document,
        "Deténgase",
        "Si el equipo indica Intel, no continúe. El constructor actual es exclusivamente ARM64.",
        kind="danger",
    )

    add_step_heading(document, 2, "Conectar la carpeta de red")
    add_numbered_actions(
        document,
        [
            "Haga clic en Finder, el icono azul con una cara que aparece en la barra inferior.",
            "En la barra superior seleccione Ir y luego Conectarse al servidor…. También puede pulsar ⌘ + K.",
            "Pegue la dirección SMB indicada abajo.",
            "Pulse Conectar.",
            "Si aparece una selección, elija Usuario registrado.",
            "Ingrese sus credenciales corporativas y pulse Conectar.",
        ],
    )
    add_command(document, "smb://ISILONSMBPROD/Gerencia_Riesgo_De_Tesoreria")
    document.add_paragraph(
        "La unidad debe aparecer en la barra lateral de Finder, dentro de Ubicaciones. "
        "Apple documenta el mismo flujo Finder → Ir → Conectarse al servidor para recursos SMB."
    )
    p = document.add_paragraph()
    add_hyperlink(
        p,
        "Referencia oficial de Apple: conexión a servidores compartidos",
        "https://support.apple.com/en-gb/guide/mac-help/mchlp1140/mac",
    )
    add_callout(
        document,
        "Si no conecta",
        "Confirme que el Mac está en la red de la oficina o en la VPN. Si el error continúa, solicite a TI validar las credenciales y el acceso SMB.",
        kind="warning",
    )

    add_step_heading(document, 3, "Copiar el paquete constructor al escritorio")
    document.add_paragraph("En la unidad de red, abra estas carpetas en orden:")
    add_command(
        document,
        "3 Jefatura de Riesgo de Mercado\nProyecto Risko\nportal\nmacos\npaquete-construccion",
    )
    document.add_paragraph("Localice el archivo:")
    add_command(document, "Preparar Instalador RISKO Mac ARM64.zip")
    add_numbered_actions(
        document,
        [
            "Arrastre el ZIP hasta el Escritorio del Mac.",
            "Espere a que termine la copia.",
            "Haga doble clic sobre el ZIP del escritorio.",
            "Abra la carpeta descomprimida y confirme que contiene las carpetas portal y Herramientas.",
        ],
    )

    add_step_heading(document, 4, "Instalar Python")
    add_numbered_actions(
        document,
        [
            "Abra Safari desde la barra inferior.",
            "Entre en la página oficial de descargas de Python para macOS.",
            "Seleccione la última versión estable de Python 3. No use alpha, beta ni release candidate.",
            "Descargue macOS 64-bit universal2 installer.",
            "Abra el archivo .pkg descargado y pulse Continuar hasta llegar a Instalar.",
            "Acepte la licencia y autorice con Touch ID o con la contraseña del Mac.",
            "Al terminar, pulse Cerrar.",
        ],
    )
    p = document.add_paragraph()
    add_hyperlink(
        p,
        "Descarga oficial: Python para macOS",
        "https://www.python.org/downloads/macos/",
    )
    add_callout(
        document,
        "Versión requerida",
        "El constructor necesita Python 3.11 o posterior con tkinter. El instalador universal2 incluye soporte para Apple Silicon.",
        kind="info",
    )

    add_step_heading(document, 5, "Abrir Terminal")
    add_numbered_actions(
        document,
        [
            "Pulse simultáneamente ⌘ + Espacio para abrir Spotlight.",
            "Escriba Terminal.",
            "Pulse Enter o Return.",
        ],
    )
    document.add_paragraph(
        "Aparecerá una ventana con una línea esperando instrucciones. Para pegar un comando use ⌘ + V y luego pulse Enter."
    )

    add_step_heading(document, 6, "Verificar arquitectura, Python y tkinter")
    document.add_paragraph("Pegue cada comando por separado y pulse Enter después de cada uno.")
    add_command(document, "uname -m", "arm64")
    add_command(document, "python3 --version", "Python 3.11 o una versión superior")
    add_command(
        document,
        "python3 -c \"import tkinter; print('tkinter OK')\"",
        "tkinter OK",
    )
    add_callout(
        document,
        "Deténgase ante un error",
        "Si uname muestra x86_64, Python no aparece o tkinter falla, no continúe. Copie el mensaje completo para solicitar soporte.",
        kind="danger",
    )

    add_step_heading(document, 7, "Instalar las herramientas de Apple")
    document.add_paragraph("En Terminal ejecute:")
    add_command(document, "xcode-select --install")
    add_numbered_actions(
        document,
        [
            "En la ventana que aparece, pulse Instalar.",
            "Acepte la licencia.",
            "Espere a que termine la descarga e instalación.",
            "Pulse Finalizar o Aceptar.",
        ],
    )
    document.add_paragraph("Después confirme la instalación:")
    add_command(document, "xcode-select -p", "/Library/Developer/CommandLineTools")
    p = document.add_paragraph()
    add_hyperlink(
        p,
        "Referencia oficial de Apple: Command Line Tools",
        "https://developer.apple.com/documentation/xcode/installing-the-command-line-tools",
    )

    add_step_heading(document, 8, "Entrar a la carpeta del constructor")
    document.add_paragraph(
        "Para evitar errores con espacios, inserte la ruta arrastrando la carpeta desde Finder."
    )
    add_numbered_actions(
        document,
        [
            "En Terminal escriba cd y deje un espacio después de la letra d. No pulse Enter todavía.",
            "En Finder abra la carpeta descomprimida del escritorio, luego portal y después macos.",
            "Arrastre la carpeta macos desde Finder hasta la ventana de Terminal.",
            "Cuando Terminal complete la ruta, pulse Enter.",
        ],
    )
    add_command(document, "pwd", "una ruta que termine en /portal/macos")

    add_step_heading(document, 9, "Habilitar el constructor")
    document.add_paragraph("Copie, pegue y ejecute:")
    add_command(
        document,
        'chmod +x construir_instalador_risko_mac.sh "Construir Instalador RISKO Mac.command"',
    )
    add_callout(
        document,
        "Comportamiento normal",
        "Si el comando no muestra ningún mensaje y vuelve a aparecer la línea de Terminal, terminó correctamente.",
        kind="success",
    )

    add_step_heading(document, 10, "Construir y publicar el instalador")
    document.add_paragraph(
        "Utilice nuevamente el método de arrastrar la carpeta para insertar la ruta productiva."
    )
    add_numbered_actions(
        document,
        [
            "En Terminal escriba el texto mostrado abajo y deje un espacio al final. No pulse Enter todavía.",
            "En Finder, dentro de la unidad de red, ubique la carpeta Portal Riesgos de Mercado.",
            "Arrastre esa carpeta hasta Terminal.",
            "Pulse Enter para comenzar.",
        ],
    )
    add_command(document, "./construir_instalador_risko_mac.sh --publish-root ")
    document.add_paragraph("El comando completo se verá aproximadamente así:")
    add_command(
        document,
        "./construir_instalador_risko_mac.sh \\\n"
        '  --publish-root "/Volumes/Gerencia_Riesgo_De_Tesoreria/Portal Riesgos de Mercado"',
    )
    add_callout(
        document,
        "Durante la construcción",
        "Aparecerán muchas líneas. No cierre Terminal, no desconecte la VPN y no permita que el Mac se suspenda. El proceso puede tardar varios minutos.",
        kind="warning",
    )
    document.add_paragraph("Al finalizar deben aparecer mensajes similares a estos:")
    add_command(
        document,
        "Instalador creado: .../Instalador RISKO Mac.pkg\n"
        "Publicado sin modificar Windows:\n"
        ".../Sistema Portal/Instalador/Instalador RISKO Mac.pkg\n"
        ".../Sistema Portal/portal-macos.json\n"
        ".../Sistema Portal/Versiones Mac/2.1.9",
    )
    add_callout(
        document,
        "Ante un error",
        "Si aparece ERROR, Permission denied, Failed o No se encontró, no cierre Terminal. Copie las últimas líneas para diagnóstico.",
        kind="danger",
    )

    add_step_heading(document, 11, "Localizar el instalador terminado")
    document.add_paragraph("En Finder abra la siguiente ubicación de red:")
    add_command(
        document,
        "Gerencia_Riesgo_De_Tesoreria\n"
        "Portal Riesgos de Mercado\n"
        "Sistema Portal\n"
        "Instalador\n"
        "Instalador RISKO Mac.pkg",
    )
    add_callout(
        document,
        "Distribución",
        "El archivo que se entrega a los usuarios es Instalador RISKO Mac.pkg. No distribuya el ZIP constructor ni los scripts.",
        kind="info",
    )

    add_step_heading(document, 12, "Instalar Portal RISKO")
    add_numbered_actions(
        document,
        [
            "Haga doble clic en Instalador RISKO Mac.pkg.",
            "Pulse Continuar.",
            "Pulse Instalar.",
            "Autorice con Touch ID o con la contraseña del Mac.",
            "Espere la confirmación de instalación correcta y pulse Cerrar.",
        ],
    )
    add_callout(
        document,
        "Primera versión piloto",
        "Si todavía no se configuró un certificado Developer ID, macOS puede advertir que no puede verificar al desarrollador. La autorización debe realizarse con aprobación de TI.",
        kind="warning",
    )
    document.add_heading("Si macOS bloquea el instalador", level=2)
    add_numbered_actions(
        document,
        [
            "Intente abrir el .pkg una vez y cierre el aviso.",
            "Abra  → Configuración del Sistema.",
            "Seleccione Privacidad y seguridad.",
            "Busque el aviso relacionado con Instalador RISKO Mac y pulse Abrir de todos modos.",
            "Autorice con Touch ID o contraseña y vuelva a abrir el .pkg.",
        ],
    )
    add_callout(
        document,
        "Mac administrado",
        "Si Abrir de todos modos no aparece o está bloqueado, TI debe autorizar la aplicación mediante las políticas corporativas.",
        kind="danger",
    )

    add_step_heading(document, 13, "Abrir Portal RISKO")
    add_numbered_actions(
        document,
        [
            "Abra Finder.",
            "Seleccione Aplicaciones en la barra lateral. También puede pulsar ⇧ + ⌘ + A.",
            "Busque Portal RISKO.",
            "Haga doble clic para abrirlo.",
        ],
    )
    document.add_paragraph(
        "Debe aparecer primero la ventana de carga y luego el portal. Como la unidad de red ya está montada, el launcher podrá consultar el release y los dashboards."
    )

    add_step_heading(document, 14, "Realizar la comprobación final")
    add_bullets(
        document,
        [
            "Portal RISKO abre desde la carpeta Aplicaciones.",
            "Se muestran los dashboards publicados.",
            "El dashboard de posición abre en el navegador.",
            "Portal RISKO se puede cerrar con ⌘ + Q.",
            "La segunda apertura es más rápida porque utiliza el caché local.",
            "La carpeta de producción contiene portal-macos.json y Versiones Mac/2.1.9.",
            "El instalador Windows y portal.json permanecen intactos.",
        ],
        checkbox=True,
    )
    document.add_paragraph("Para abrir la carpeta de diagnóstico ejecute:")
    add_command(
        document,
        'open "$HOME/Library/Application Support/PortalRisko/Launcher"',
    )
    document.add_paragraph("El archivo principal de diagnóstico es:")
    add_command(document, "launcher-macos.log")

    document.add_page_break()
    document.add_heading("Solución de problemas", level=1)
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Situación o mensaje", "Qué significa", "Qué hacer")
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, PRIMARY)
        set_cell_border(cell)
        set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    issues = (
        (
            "El servidor no conecta",
            "El Mac no ve la carpeta SMB.",
            "Conectarse a la red de oficina o VPN. Reintentar con ⌘ + K y validar credenciales con TI.",
        ),
        (
            "uname muestra x86_64",
            "Terminal está ejecutándose mediante Rosetta o el equipo no es ARM.",
            "Confirmar el chip en Acerca de este Mac. En un Apple Silicon, pedir a TI desactivar Rosetta para Terminal.",
        ),
        (
            "python3: command not found",
            "Python no quedó instalado o Terminal no se reinició.",
            "Cerrar y abrir Terminal. Si continúa, reinstalar el paquete universal2 oficial de Python.",
        ),
        (
            "No module named tkinter",
            "El Python instalado no incluye la interfaz gráfica.",
            "Instalar Python desde python.org usando el instalador universal2.",
        ),
        (
            "xcode-select solicita instalación",
            "Faltan las herramientas de Apple.",
            "Aceptar la instalación, esperar a que termine y ejecutar nuevamente xcode-select -p.",
        ),
        (
            "Permission denied",
            "Faltan permisos de ejecución o escritura.",
            "Ejecutar nuevamente chmod. Si ocurre al publicar en red, solicitar a TI permiso de escritura.",
        ),
        (
            "pip / SSL / proxy / timeout",
            "El Mac no puede descargar dependencias.",
            "Confirmar internet y proxy corporativo. Entregar el mensaje completo a TI.",
        ),
        (
            "La versión Mac 2.1.9 ya existe",
            "El release ya fue publicado y es inmutable.",
            "No borrar ni sobrescribir. Verificar si el instalador ya está disponible o publicar una nueva versión del código.",
        ),
        (
            "Apple no puede verificar al desarrollador",
            "El piloto no está notarizado con Developer ID.",
            "Solicitar autorización a TI y usar Privacidad y seguridad → Abrir de todos modos.",
        ),
        (
            "Portal abre sin dashboards",
            "La carpeta productiva no está montada o no responde.",
            "Reconectar el SMB, abrir nuevamente Portal RISKO y revisar launcher-macos.log.",
        ),
    )
    for values in issues:
        row = table.add_row()
        prevent_row_split(row)
        for cell, value in zip(row.cells, values):
            set_cell_border(cell)
            set_cell_margins(cell)
            set_cell_shading(cell, "FFFFFF")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.paragraphs[0].add_run(value)

    document.add_heading("Información que debe enviarse al solicitar ayuda", level=2)
    add_bullets(
        document,
        [
            "Fotografía o captura del mensaje de error.",
            "Las últimas 20 a 30 líneas mostradas en Terminal.",
            "Resultado de uname -m.",
            "Resultado de python3 --version.",
            "Archivo launcher-macos.log si Portal RISKO alcanzó a instalarse.",
        ],
        checkbox=True,
    )

    document.add_heading("Referencias", level=1)
    references = (
        (
            "Apple — Conectarse a computadores y servidores compartidos",
            "https://support.apple.com/en-gb/guide/mac-help/mchlp1140/mac",
        ),
        (
            "Apple Developer — Instalar Command Line Tools",
            "https://developer.apple.com/documentation/xcode/installing-the-command-line-tools",
        ),
        (
            "Python — Descargas oficiales para macOS",
            "https://www.python.org/downloads/macos/",
        ),
    )
    for label, url in references:
        p = document.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)

    add_callout(
        document,
        "Cierre",
        "Una vez validado el piloto, distribuya únicamente Instalador RISKO Mac.pkg. La firma y notarización corporativa deben completarse antes de una distribución general sin advertencias de seguridad.",
        kind="info",
    )
    return document


def main() -> None:
    document = build_document()
    document.core_properties.title = (
        "Instructivo paso a paso - Construcción Instalador RISKO Mac"
    )
    document.core_properties.subject = "Portal RISKO para macOS ARM64"
    document.core_properties.author = "Riesgo de Mercado"
    document.core_properties.keywords = "RISKO, macOS, ARM64, instalador, Portal"
    document.core_properties.comments = "Documento de uso interno"
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
