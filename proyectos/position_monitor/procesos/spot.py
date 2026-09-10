"""Posicion Spot acumulada a partir del Reporte de Caja.

La fuente de Summit tiene extension ``.xls``, pero internamente es texto con
separador punto y coma. El calculo monetario se realiza con ``Decimal`` y se
persiste como texto decimal en SQLite para no introducir redondeos binarios.

Ademas de los owner tables generales, se incorporan los colaterales DP de
``FWD_CLIENTES`` con el mismo filtro del libro operativo ReporteCaja(con DP).
"""

from __future__ import annotations

import csv
from contextlib import closing
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import getpass
import hashlib
import json
from pathlib import Path
import sys
import uuid

import pandas as pd


RAIZ_RISKO = Path(__file__).resolve().parents[3]
if str(RAIZ_RISKO) not in sys.path:
    sys.path.insert(0, str(RAIZ_RISKO))

from compartido.nucleo_risko.archivos import registrar_log
from compartido.nucleo_risko.base_datos import conectar_sqlite
from compartido.nucleo_risko.tablas_auxiliares import guardar_tablas_auxiliares_modulo
from compartido.nucleo_risko.rutas import (
    ARCHIVO_BASE_DATOS_POSITION_MONITOR,
    ARCHIVO_BASE_DATOS_RISKO_PRUEBAS,
    ARCHIVO_CONFIG_RISKO_VECTOR,
    CARPETA_INSUMOS,
)
from proyectos.position_monitor.procesos.parametros_libros import (
    cargar_parametros_libros,
    normalizar_clave_param,
)


TABLA_HISTORICO_SPOT = "tbl_posicion_spot_acumulada"
TABLA_MOVIMIENTOS_SPOT = "tbl_movimientos_spot"
TABLA_ALERTAS_SPOT = "tbl_alertas_spot"
TABLA_EJECUCIONES_SPOT = "tbl_ejecuciones_spot"
TABLA_AUDITORIA_SPOT = "tbl_auditoria_spot"
TABLA_COMPRAS_VENTAS_BOOK_SPOT = "tbl_compras_ventas_spot_book"

ORIGEN_MOVIMIENTOS = "MOVIMIENTOS"
ORIGEN_CARRY = "CARRY_FORWARD"
ORIGEN_EXCEL = "EXCEL_INICIAL"
ORIGEN_INICIAL_CERO = "INICIAL_CERO"

CODIGOS_ALERTA = {
    "UNKNOWN_BOOK",
    "INACTIVE_BOOK",
    "MISSING_PREVIOUS_POSITION",
    "DUPLICATE_SOURCE_FILE",
    "SOURCE_FILE_CHANGE",
    "INVALID_AMOUNT",
    "INVALID_DATE",
    "HISTORICAL_REPROCESSING",
    "INITIAL_POSITION_CHANGE",
}

OWNER_TABLES_POR_DEFECTO = {
    "CONTADO",
    "FXFWD",
    "FXOPT_TR",
    "FXSPOT",
    "MM",
    "NOVADO",
    "SWAP",
}

VERSION_LOGICA_SPOT = "2026-08-20-SPOT-CLIENTE-TRADE-DATE-V2"
OWNER_COLATERALES_DP = "DPMT_TR"
BOOK_COLATERALES_DP = "FWD_CLIENTES"
BOOK_FECHA_TRADE_DATE = "SPOT_CLIENTE"
EVTYPES_COLATERALES_DP = {"FEE", "INT"}

COLUMNAS_REQUERIDAS = (
    "GeneratedPK",
    "DmOwnerTable",
    "TradeId",
    "Book",
    "Desk",
    "EvType",
    "SettleCcy",
    "Amount",
    "ValueDate",
    "TRADE DATE",
)

COLUMNAS_POSICION = [
    "FECHA",
    "PRODUCTO",
    "BOOK",
    "POSICION",
    "MONEDA_POSICION",
    "LB_LT",
    "INSTRUMENTO",
    "COMPANY",
    "CLASIFICACION_CONTABLE",
    "BANKING_CVA_DVA",
]

_ACTIVOS = {"SI", "S", "TRUE", "1", "ACTIVO", "ACTIVA", "YES"}
_INACTIVOS = {"NO", "N", "FALSE", "0", "INACTIVO", "INACTIVA"}


def _ahora() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _usuario() -> str:
    try:
        return getpass.getuser() or "DESCONOCIDO"
    except Exception:
        return "DESCONOCIDO"


def _fecha(valor: object) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha no valida para Spot: {valor}")


def _fecha_texto(valor: object) -> str:
    return _fecha(valor).strftime("%d/%m/%Y")


def _decimal(valor: object) -> Decimal:
    texto = str(valor).strip().replace("\u00a0", "").replace(" ", "")
    if texto.startswith("(") and texto.endswith(")"):
        texto = f"-{texto[1:-1]}"
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            # Formato local: 50.269.423,25732988
            texto = texto.replace(".", "").replace(",", ".")
        else:
            # Formato Summit: 50,269,423.25732988
            texto = texto.replace(",", "")
    elif "," in texto:
        partes = texto.split(",")
        if len(partes) == 2 and len(partes[1]) != 3:
            texto = ".".join(partes)
        else:
            texto = "".join(partes)
    if not texto:
        raise InvalidOperation("valor vacio")
    resultado = Decimal(texto)
    if not resultado.is_finite():
        raise InvalidOperation("valor no finito")
    return resultado


def _decimal_texto(valor: Decimal | object) -> str:
    numero = valor if isinstance(valor, Decimal) else _decimal(valor)
    if numero == 0:
        return "0"
    return format(numero, "f")


def _cargar_configuracion_spot() -> dict:
    configuracion: dict = {}
    if ARCHIVO_CONFIG_RISKO_VECTOR.exists():
        with ARCHIVO_CONFIG_RISKO_VECTOR.open("r", encoding="utf-8") as archivo:
            raiz = json.load(archivo)
        configuracion = dict(raiz.get("spot") or {})

    owners = configuracion.get("owner_tables_incluidas") or sorted(
        OWNER_TABLES_POR_DEFECTO
    )
    configuracion["owner_tables_incluidas"] = {
        str(owner).strip().upper() for owner in owners if str(owner).strip()
    }
    if not configuracion["owner_tables_incluidas"]:
        raise ValueError("La configuracion Spot no tiene owner tables incluidas.")
    configuracion.setdefault(
        "archivo_excel_posiciones",
        str(
            RAIZ_RISKO
            / "proyectos"
            / "position_monitor"
            / "configuracion"
            / "param_posiciones_iniciales_spot.xlsx"
        ),
    )
    configuracion.setdefault("hoja_excel_posiciones", "Posiciones_Iniciales")
    return configuracion


def _resolver_ruta_db(
    ruta_db: str | Path | None = None,
    pruebas: bool = False,
) -> Path:
    if ruta_db is not None:
        return Path(ruta_db)
    if pruebas:
        return ARCHIVO_BASE_DATOS_RISKO_PRUEBAS
    return ARCHIVO_BASE_DATOS_POSITION_MONITOR


def _ruta_insumo(fecha_corte: date) -> Path:
    return CARPETA_INSUMOS / f"Cierre_ReporteCaja_{fecha_corte:%d%m%y}_000.xls"


def _hash_archivo(ruta: Path) -> str:
    digest = hashlib.sha256()
    with ruta.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def _valor_activo(valor: object) -> bool:
    clave = normalizar_clave_param(valor)
    if clave in _ACTIVOS:
        return True
    if clave in _INACTIVOS:
        return False
    raise ValueError(
        "Valor ACTIVO invalido en param_libros.csv para Spot: "
        f"{valor!r}. Use SI o NO."
    )


def _cargar_parametros_spot() -> tuple[dict[str, dict], dict[str, dict], str]:
    parametros = cargar_parametros_libros()
    parametros = parametros.loc[parametros["_PRODUCTO_CLAVE"] == "SPOT"].copy()
    if parametros.empty:
        raise ValueError(
            "param_libros.csv no contiene definiciones para PRODUCTO = Spot."
        )

    por_book: dict[str, dict] = {}
    por_posicion: dict[str, dict] = {}
    serializables: list[dict] = []
    for _, fila in parametros.iterrows():
        book = str(fila["BOOK"]).strip()
        book_clave = normalizar_clave_param(book)
        nombre = str(fila["BOOK_HOMOL"]).strip() or book
        posicion_id = normalizar_clave_param(nombre)
        if not book_clave or not posicion_id:
            raise ValueError("Spot tiene un BOOK o BOOK_HOMOL vacio en param_libros.csv.")
        activo = _valor_activo(fila.get("ACTIVO", "SI"))
        definicion = {
            "BOOK": book,
            "BOOK_CLAVE": book_clave,
            "POSICION_ID": posicion_id,
            "POSICION_NOMBRE": nombre,
            "ACTIVO": activo,
            "LB_LT": str(fila["LB_LT"]).strip() or "No definido",
            "INSTRUMENTO": str(fila["INSTRUMENTO"]).strip() or "Derivados",
            "MONEDA_POSICION": str(fila["MONEDA_POSICION"]).strip() or "USD",
        }
        por_book[book_clave] = definicion

        existente = por_posicion.get(posicion_id)
        if existente:
            campos = ("POSICION_NOMBRE", "ACTIVO", "LB_LT", "INSTRUMENTO", "MONEDA_POSICION")
            diferentes = [campo for campo in campos if existente[campo] != definicion[campo]]
            if diferentes:
                raise ValueError(
                    "param_libros.csv tiene aliases Spot incompatibles para la "
                    f"posicion {nombre}: {', '.join(diferentes)}."
                )
        else:
            por_posicion[posicion_id] = definicion
        serializables.append(definicion)

    huella = hashlib.sha256(
        json.dumps(
            sorted(serializables, key=lambda item: item["BOOK_CLAVE"]),
            ensure_ascii=False,
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()
    return por_book, por_posicion, huella


def _detectar_encabezado(ruta: Path) -> tuple[str, int, list[str], dict[str, int]]:
    errores: list[str] = []
    requeridas = {nombre.casefold() for nombre in COLUMNAS_REQUERIDAS}
    for codificacion in ("cp1252", "latin1"):
        try:
            with ruta.open("r", encoding=codificacion, errors="strict", newline="") as archivo:
                for indice, fila in enumerate(csv.reader(archivo, delimiter=";")):
                    if indice >= 200:
                        break
                    nombres = [str(valor).strip() for valor in fila]
                    mapa = {nombre.casefold(): pos for pos, nombre in enumerate(nombres)}
                    if requeridas.issubset(mapa):
                        return codificacion, indice, nombres, mapa
        except UnicodeDecodeError as exc:
            errores.append(f"{codificacion}: {exc}")
    raise ValueError(
        "No se encontro el encabezado requerido del Reporte de Caja en las "
        "primeras 200 filas. Se exigen ValueDate y TRADE DATE, junto con las columnas "
        f"{', '.join(COLUMNAS_REQUERIDAS)}. Detalle: {'; '.join(errores)}"
    )


def _alerta(
    codigo: str,
    fecha_corte: str,
    detalle: str,
    *,
    posicion_id: str = "",
    book_original: str = "",
    fila_fuente: int | None = None,
) -> dict:
    if codigo not in CODIGOS_ALERTA:
        raise ValueError(f"Codigo de alerta Spot no reconocido: {codigo}")
    return {
        "CODIGO": codigo,
        "FECHA": fecha_corte,
        "POSICION_ID": posicion_id,
        "BOOK_ORIGINAL": book_original,
        "FILA_FUENTE": fila_fuente,
        "DETALLE": detalle,
    }


def _es_colateral_dp(
    owner: str,
    book: str,
    evtype: str,
    settle_ccy: str,
) -> bool:
    """Replica el filtro de colaterales de ReporteCaja(con DP).xlsb."""
    return (
        str(owner).strip().upper() == OWNER_COLATERALES_DP
        and normalizar_clave_param(book) == BOOK_COLATERALES_DP
        and str(evtype).strip().upper() in EVTYPES_COLATERALES_DP
        and str(settle_ccy).strip().upper() == "USD"
    )


def _leer_movimientos(
    ruta: Path,
    fecha_corte: date,
    por_book: dict[str, dict],
    owners_incluidos: set[str],
    logger=None,
) -> tuple[list[dict], list[dict]]:
    codificacion, fila_encabezado, _, mapa = _detectar_encabezado(ruta)
    fecha_objetivo = fecha_corte.strftime("%d/%m/%Y")
    registrar_log(
        logger,
        f"Spot: encabezado detectado en fila {fila_encabezado + 1}, "
        f"codificacion {codificacion}.",
    )

    movimientos: list[dict] = []
    alertas: list[dict] = []
    alertados: set[tuple[str, str]] = set()
    cantidad_colaterales_dp = 0
    neto_colaterales_dp = Decimal("0")

    def campo(fila: list[str], nombre: str) -> str:
        indice = mapa[nombre.casefold()]
        return str(fila[indice]).strip() if indice < len(fila) else ""

    with ruta.open("r", encoding=codificacion, errors="strict", newline="") as archivo:
        lector = csv.reader(archivo, delimiter=";")
        for _ in range(fila_encabezado + 1):
            next(lector, None)

        for numero_fila, fila in enumerate(lector, start=fila_encabezado + 2):
            owner = campo(fila, "DmOwnerTable").upper()
            book_original = campo(fila, "Book")
            book_clave = normalizar_clave_param(book_original)
            evtype = campo(fila, "EvType").upper()
            moneda = campo(fila, "SettleCcy").upper()
            es_colateral_dp = _es_colateral_dp(
                owner,
                book_original,
                evtype,
                moneda,
            )
            if owner not in owners_incluidos and not es_colateral_dp:
                continue
            if moneda != "USD":
                continue

            columna_fecha = (
                "TRADE DATE" if book_clave == BOOK_FECHA_TRADE_DATE else "ValueDate"
            )
            fecha_original = campo(fila, columna_fecha)
            try:
                fecha_movimiento = _fecha(fecha_original)
            except ValueError:
                alertas.append(
                    _alerta(
                        "INVALID_DATE",
                        fecha_objetivo,
                        f"{columna_fecha} invalida: {fecha_original!r}.",
                        book_original=campo(fila, "Book"),
                        fila_fuente=numero_fila,
                    )
                )
                continue
            if fecha_movimiento != fecha_corte:
                continue

            definicion = por_book.get(book_clave)
            posicion_id = definicion["POSICION_ID"] if definicion else ""
            posicion_nombre = definicion["POSICION_NOMBRE"] if definicion else ""
            publicable = bool(definicion and definicion["ACTIVO"])

            if definicion is None:
                llave = ("UNKNOWN_BOOK", book_clave)
                if llave not in alertados:
                    alertados.add(llave)
                    alertas.append(
                        _alerta(
                            "UNKNOWN_BOOK",
                            fecha_objetivo,
                            "El BOOK no esta definido para Spot en param_libros.csv; "
                            "sus movimientos no se agregan a la posicion oficial.",
                            book_original=book_original,
                        )
                    )
            elif not definicion["ACTIVO"]:
                llave = ("INACTIVE_BOOK", book_clave)
                if llave not in alertados:
                    alertados.add(llave)
                    alertas.append(
                        _alerta(
                            "INACTIVE_BOOK",
                            fecha_objetivo,
                            "El BOOK esta inactivo en param_libros.csv; sus movimientos "
                            "no se agregan a la posicion oficial.",
                            posicion_id=posicion_id,
                            book_original=book_original,
                        )
                    )

            amount_original = campo(fila, "Amount")
            try:
                amount = _decimal(amount_original)
            except (InvalidOperation, ValueError):
                alertas.append(
                    _alerta(
                        "INVALID_AMOUNT",
                        fecha_objetivo,
                        f"Amount invalido: {amount_original!r}.",
                        posicion_id=posicion_id,
                        book_original=book_original,
                        fila_fuente=numero_fila,
                    )
                )
                continue

            if es_colateral_dp:
                cantidad_colaterales_dp += 1
                neto_colaterales_dp += amount

            if amount > 0:
                tipo = "COMPRA"
                monto_usd = amount
            elif amount < 0:
                tipo = "VENTA"
                monto_usd = -amount
            else:
                tipo = "SIN_MOVIMIENTO"
                monto_usd = Decimal("0")

            movimientos.append(
                {
                    "FECHA": fecha_objetivo,
                    "SOURCE_ROW_ID": campo(fila, "GeneratedPK") or str(numero_fila),
                    "TRADE_ID": campo(fila, "TradeId"),
                    "PRODUCTO_ORIGEN": owner,
                    "BOOK_ORIGINAL": book_original,
                    "POSICION_ID": posicion_id,
                    "POSICION_NOMBRE": posicion_nombre,
                    "DESK": campo(fila, "Desk"),
                    "EVTYPE": evtype,
                    "SETTLE_CCY": moneda,
                    "AMOUNT_ORIGINAL": _decimal_texto(amount),
                    "TIPO_MOVIMIENTO": tipo,
                    "MONTO_USD": _decimal_texto(monto_usd),
                    "FILA_FUENTE": numero_fila,
                    "PUBLICABLE": 1 if publicable else 0,
                }
            )

    registrar_log(
        logger,
        f"Spot: {len(movimientos):,} movimientos USD del corte y "
        f"{len(alertas):,} alertas de lectura/mapeo. "
        f"Colaterales DP FWD_CLIENTES incluidos: {cantidad_colaterales_dp:,}, "
        f"neto USD {_decimal_texto(neto_colaterales_dp)}.",
    )
    return movimientos, alertas


def _crear_esquema(conexion) -> None:
    conexion.executescript(
        f"""
        CREATE TABLE IF NOT EXISTS {TABLA_EJECUCIONES_SPOT} (
            EJECUCION_ID TEXT PRIMARY KEY,
            FECHA_CORTE TEXT NOT NULL,
            ARCHIVO_FUENTE TEXT NOT NULL,
            RUTA_FUENTE TEXT NOT NULL,
            SHA256_FUENTE TEXT NOT NULL,
            SHA256_PARAMETROS TEXT NOT NULL,
            TAMANO_BYTES INTEGER NOT NULL,
            MODIFICADO_FUENTE TEXT,
            ORIGEN TEXT NOT NULL,
            ESTADO TEXT NOT NULL,
            USUARIO TEXT NOT NULL,
            INICIO TEXT NOT NULL,
            FIN TEXT,
            MENSAJE TEXT
        );

        CREATE TABLE IF NOT EXISTS {TABLA_MOVIMIENTOS_SPOT} (
            MOVIMIENTO_ID TEXT PRIMARY KEY,
            EJECUCION_ID TEXT NOT NULL,
            FECHA TEXT NOT NULL,
            SOURCE_ROW_ID TEXT,
            TRADE_ID TEXT,
            PRODUCTO_ORIGEN TEXT,
            BOOK_ORIGINAL TEXT,
            POSICION_ID TEXT,
            POSICION_NOMBRE TEXT,
            DESK TEXT,
            EVTYPE TEXT,
            SETTLE_CCY TEXT,
            AMOUNT_ORIGINAL TEXT,
            TIPO_MOVIMIENTO TEXT,
            MONTO_USD TEXT,
            ARCHIVO_FUENTE TEXT,
            FILA_FUENTE INTEGER,
            PUBLICABLE INTEGER NOT NULL,
            VIGENTE INTEGER NOT NULL,
            CREATED_AT TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS {TABLA_COMPRAS_VENTAS_BOOK_SPOT} (
            FECHA TEXT NOT NULL,
            BOOK_CLAVE TEXT NOT NULL,
            BOOK_ORIGINAL TEXT NOT NULL,
            POSICION_ID TEXT,
            POSICION_NOMBRE TEXT,
            ESTADO_MAPEO TEXT NOT NULL,
            PUBLICABLE INTEGER NOT NULL,
            COMPRAS_USD TEXT NOT NULL,
            VENTAS_USD TEXT NOT NULL,
            MOVIMIENTO_NETO_USD TEXT NOT NULL,
            NUMERO_MOVIMIENTOS INTEGER NOT NULL,
            ARCHIVO_FUENTE TEXT NOT NULL,
            EJECUCION_ID TEXT NOT NULL,
            UPDATED_AT TEXT NOT NULL,
            PRIMARY KEY (FECHA, BOOK_CLAVE)
        );

        CREATE TABLE IF NOT EXISTS {TABLA_ALERTAS_SPOT} (
            ALERTA_ID TEXT PRIMARY KEY,
            EJECUCION_ID TEXT NOT NULL,
            FECHA TEXT NOT NULL,
            CODIGO TEXT NOT NULL,
            POSICION_ID TEXT,
            BOOK_ORIGINAL TEXT,
            FILA_FUENTE INTEGER,
            DETALLE TEXT NOT NULL,
            VIGENTE INTEGER NOT NULL,
            CREATED_AT TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS {TABLA_HISTORICO_SPOT} (
            FECHA TEXT NOT NULL,
            POSICION_ID TEXT NOT NULL,
            POSICION_NOMBRE TEXT NOT NULL,
            FECHA_ANTERIOR TEXT,
            POSICION_ANTERIOR TEXT NOT NULL,
            COMPRAS_USD TEXT NOT NULL,
            VENTAS_USD TEXT NOT NULL,
            MOVIMIENTO_NETO_USD TEXT NOT NULL,
            POSICION_FINAL_USD TEXT NOT NULL,
            ARCHIVO_FUENTE TEXT,
            EJECUCION_ID TEXT NOT NULL,
            TIPO_ORIGEN TEXT NOT NULL,
            ESTADO TEXT NOT NULL,
            TIENE_ALERTA INTEGER NOT NULL,
            DETALLE_ALERTA TEXT,
            USUARIO TEXT NOT NULL,
            CREATED_AT TEXT NOT NULL,
            UPDATED_AT TEXT NOT NULL,
            PRIMARY KEY (FECHA, POSICION_ID)
        );

        CREATE TABLE IF NOT EXISTS {TABLA_AUDITORIA_SPOT} (
            AUDITORIA_ID TEXT PRIMARY KEY,
            EJECUCION_ID TEXT NOT NULL,
            FECHA TEXT NOT NULL,
            POSICION_ID TEXT NOT NULL,
            VALOR_ANTERIOR TEXT,
            VALOR_NUEVO TEXT,
            MOTIVO TEXT NOT NULL,
            CREATED_AT TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_mov_spot_fecha_pos
            ON {TABLA_MOVIMIENTOS_SPOT} (FECHA, POSICION_ID, VIGENTE);
        CREATE INDEX IF NOT EXISTS idx_book_spot_fecha_publicable
            ON {TABLA_COMPRAS_VENTAS_BOOK_SPOT} (FECHA, PUBLICABLE);
        CREATE INDEX IF NOT EXISTS idx_alerta_spot_fecha
            ON {TABLA_ALERTAS_SPOT} (FECHA, POSICION_ID, VIGENTE);
        CREATE INDEX IF NOT EXISTS idx_ejec_spot_fuente
            ON {TABLA_EJECUCIONES_SPOT}
            (FECHA_CORTE, SHA256_FUENTE, SHA256_PARAMETROS, ESTADO);
        """
    )


def _insertar_alerta(conexion, ejecucion_id: str, datos: dict) -> None:
    conexion.execute(
        f"""
        INSERT INTO {TABLA_ALERTAS_SPOT} (
            ALERTA_ID, EJECUCION_ID, FECHA, CODIGO, POSICION_ID,
            BOOK_ORIGINAL, FILA_FUENTE, DETALLE, VIGENTE, CREATED_AT
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
        """,
        (
            str(uuid.uuid4()),
            ejecucion_id,
            datos["FECHA"],
            datos["CODIGO"],
            datos.get("POSICION_ID", ""),
            datos.get("BOOK_ORIGINAL", ""),
            datos.get("FILA_FUENTE"),
            datos["DETALLE"],
            _ahora(),
        ),
    )


def _insertar_auditoria(
    conexion,
    ejecucion_id: str,
    fecha_texto: str,
    posicion_id: str,
    anterior: str | None,
    nuevo: str | None,
    motivo: str,
) -> None:
    conexion.execute(
        f"""
        INSERT INTO {TABLA_AUDITORIA_SPOT} (
            AUDITORIA_ID, EJECUCION_ID, FECHA, POSICION_ID,
            VALOR_ANTERIOR, VALOR_NUEVO, MOTIVO, CREATED_AT
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            str(uuid.uuid4()),
            ejecucion_id,
            fecha_texto,
            posicion_id,
            anterior,
            nuevo,
            motivo,
            _ahora(),
        ),
    )


def _filas_historico_posicion(conexion, posicion_id: str) -> list[dict]:
    cursor = conexion.execute(
        f"SELECT * FROM {TABLA_HISTORICO_SPOT} WHERE POSICION_ID = ?",
        (posicion_id,),
    )
    columnas = [columna[0] for columna in cursor.description]
    return [dict(zip(columnas, fila)) for fila in cursor.fetchall()]


def _desactivar_alertas_desde(
    conexion,
    posicion_id: str,
    fecha_inicio: date,
    codigos: tuple[str, ...] = ("MISSING_PREVIOUS_POSITION",),
) -> None:
    marcadores = ",".join("?" for _ in codigos)
    filas = conexion.execute(
        f"""
        SELECT ALERTA_ID, FECHA FROM {TABLA_ALERTAS_SPOT}
        WHERE POSICION_ID = ? AND VIGENTE = 1 AND CODIGO IN ({marcadores})
        """,
        (posicion_id, *codigos),
    ).fetchall()
    ids = [alerta_id for alerta_id, fecha_texto in filas if _fecha(fecha_texto) >= fecha_inicio]
    if ids:
        conexion.executemany(
            f"UPDATE {TABLA_ALERTAS_SPOT} SET VIGENTE = 0 WHERE ALERTA_ID = ?",
            [(alerta_id,) for alerta_id in ids],
        )


def _movimientos_posicion(conexion, posicion_id: str) -> dict[date, dict]:
    filas = conexion.execute(
        f"""
        SELECT FECHA, TIPO_MOVIMIENTO, MONTO_USD, ARCHIVO_FUENTE, EJECUCION_ID
        FROM {TABLA_MOVIMIENTOS_SPOT}
        WHERE POSICION_ID = ? AND PUBLICABLE = 1 AND VIGENTE = 1
        """,
        (posicion_id,),
    ).fetchall()
    acumulados: dict[date, dict] = {}
    for fecha_texto, tipo, monto_texto, archivo, ejecucion_id in filas:
        fecha_mov = _fecha(fecha_texto)
        datos = acumulados.setdefault(
            fecha_mov,
            {
                "COMPRAS": Decimal("0"),
                "VENTAS": Decimal("0"),
                "FILAS": 0,
                "ARCHIVO": archivo,
                "EJECUCION_ID": ejecucion_id,
            },
        )
        monto = _decimal(monto_texto)
        if tipo == "COMPRA":
            datos["COMPRAS"] += monto
        elif tipo == "VENTA":
            datos["VENTAS"] += monto
        datos["FILAS"] += 1
        datos["ARCHIVO"] = archivo
        datos["EJECUCION_ID"] = ejecucion_id
    return acumulados


def _guardar_compras_ventas_por_book(
    conexion,
    fecha_texto: str,
    por_book: dict[str, dict],
    archivo_fuente: str,
    ejecucion_id: str,
) -> None:
    """Materializa compras/ventas por BOOK sin afectar la posicion canonica."""
    resumen: dict[str, dict] = {}

    # Todos los BOOK parametrizados aparecen, aunque no tengan movimientos.
    for book_clave, definicion in por_book.items():
        resumen[book_clave] = {
            "BOOK_CLAVE": book_clave,
            "BOOK_ORIGINAL": definicion["BOOK"],
            "POSICION_ID": definicion["POSICION_ID"],
            "POSICION_NOMBRE": definicion["POSICION_NOMBRE"],
            "ESTADO_MAPEO": "ACTIVO" if definicion["ACTIVO"] else "INACTIVO",
            "PUBLICABLE": 1 if definicion["ACTIVO"] else 0,
            "COMPRAS": Decimal("0"),
            "VENTAS": Decimal("0"),
            "NUMERO_MOVIMIENTOS": 0,
        }

    movimientos = conexion.execute(
        f"""
        SELECT BOOK_ORIGINAL, TIPO_MOVIMIENTO, MONTO_USD
        FROM {TABLA_MOVIMIENTOS_SPOT}
        WHERE FECHA = ? AND VIGENTE = 1
        """,
        (fecha_texto,),
    ).fetchall()
    for book_original, tipo, monto_texto in movimientos:
        book_original = str(book_original or "").strip()
        book_clave = normalizar_clave_param(book_original)
        if not book_clave:
            book_clave = "BOOK_VACIO"
        datos = resumen.get(book_clave)
        if datos is None:
            datos = {
                "BOOK_CLAVE": book_clave,
                "BOOK_ORIGINAL": book_original or "(BOOK VACIO)",
                "POSICION_ID": "",
                "POSICION_NOMBRE": "",
                "ESTADO_MAPEO": "UNKNOWN_BOOK",
                "PUBLICABLE": 0,
                "COMPRAS": Decimal("0"),
                "VENTAS": Decimal("0"),
                "NUMERO_MOVIMIENTOS": 0,
            }
            resumen[book_clave] = datos
        monto = _decimal(monto_texto)
        if tipo == "COMPRA":
            datos["COMPRAS"] += monto
        elif tipo == "VENTA":
            datos["VENTAS"] += monto
        datos["NUMERO_MOVIMIENTOS"] += 1

    conexion.execute(
        f"DELETE FROM {TABLA_COMPRAS_VENTAS_BOOK_SPOT} WHERE FECHA = ?",
        (fecha_texto,),
    )
    actualizado = _ahora()
    conexion.executemany(
        f"""
        INSERT INTO {TABLA_COMPRAS_VENTAS_BOOK_SPOT} (
            FECHA, BOOK_CLAVE, BOOK_ORIGINAL, POSICION_ID, POSICION_NOMBRE,
            ESTADO_MAPEO, PUBLICABLE, COMPRAS_USD, VENTAS_USD,
            MOVIMIENTO_NETO_USD, NUMERO_MOVIMIENTOS, ARCHIVO_FUENTE,
            EJECUCION_ID, UPDATED_AT
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                fecha_texto,
                datos["BOOK_CLAVE"],
                datos["BOOK_ORIGINAL"],
                datos["POSICION_ID"],
                datos["POSICION_NOMBRE"],
                datos["ESTADO_MAPEO"],
                datos["PUBLICABLE"],
                _decimal_texto(datos["COMPRAS"]),
                _decimal_texto(datos["VENTAS"]),
                _decimal_texto(datos["COMPRAS"] - datos["VENTAS"]),
                datos["NUMERO_MOVIMIENTOS"],
                archivo_fuente,
                ejecucion_id,
                actualizado,
            )
            for _, datos in sorted(resumen.items())
        ],
    )


def _upsert_historico(conexion, fila: dict) -> None:
    conexion.execute(
        f"""
        INSERT INTO {TABLA_HISTORICO_SPOT} (
            FECHA, POSICION_ID, POSICION_NOMBRE, FECHA_ANTERIOR,
            POSICION_ANTERIOR, COMPRAS_USD, VENTAS_USD,
            MOVIMIENTO_NETO_USD, POSICION_FINAL_USD, ARCHIVO_FUENTE,
            EJECUCION_ID, TIPO_ORIGEN, ESTADO, TIENE_ALERTA,
            DETALLE_ALERTA, USUARIO, CREATED_AT, UPDATED_AT
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(FECHA, POSICION_ID) DO UPDATE SET
            POSICION_NOMBRE = excluded.POSICION_NOMBRE,
            FECHA_ANTERIOR = excluded.FECHA_ANTERIOR,
            POSICION_ANTERIOR = excluded.POSICION_ANTERIOR,
            COMPRAS_USD = excluded.COMPRAS_USD,
            VENTAS_USD = excluded.VENTAS_USD,
            MOVIMIENTO_NETO_USD = excluded.MOVIMIENTO_NETO_USD,
            POSICION_FINAL_USD = excluded.POSICION_FINAL_USD,
            ARCHIVO_FUENTE = excluded.ARCHIVO_FUENTE,
            EJECUCION_ID = excluded.EJECUCION_ID,
            TIPO_ORIGEN = excluded.TIPO_ORIGEN,
            ESTADO = excluded.ESTADO,
            TIENE_ALERTA = excluded.TIENE_ALERTA,
            DETALLE_ALERTA = excluded.DETALLE_ALERTA,
            USUARIO = excluded.USUARIO,
            UPDATED_AT = excluded.UPDATED_AT
        """,
        tuple(
            fila[columna]
            for columna in (
                "FECHA",
                "POSICION_ID",
                "POSICION_NOMBRE",
                "FECHA_ANTERIOR",
                "POSICION_ANTERIOR",
                "COMPRAS_USD",
                "VENTAS_USD",
                "MOVIMIENTO_NETO_USD",
                "POSICION_FINAL_USD",
                "ARCHIVO_FUENTE",
                "EJECUCION_ID",
                "TIPO_ORIGEN",
                "ESTADO",
                "TIENE_ALERTA",
                "DETALLE_ALERTA",
                "USUARIO",
                "CREATED_AT",
                "UPDATED_AT",
            )
        ),
    )


def _recalcular_posicion(
    conexion,
    definicion: dict,
    fecha_inicio: date,
    fecha_minima_fin: date,
    ejecucion_id: str,
    motivo: str,
) -> None:
    posicion_id = definicion["POSICION_ID"]
    filas_anteriores = _filas_historico_posicion(conexion, posicion_id)
    por_fecha_anterior = {_fecha(fila["FECHA"]): fila for fila in filas_anteriores}
    anclas = {
        fecha_fila: fila
        for fecha_fila, fila in por_fecha_anterior.items()
        if fila["TIPO_ORIGEN"] == ORIGEN_EXCEL
    }
    movimientos = _movimientos_posicion(conexion, posicion_id)

    fechas_relevantes = [fecha_minima_fin]
    fechas_relevantes.extend(fecha for fecha in por_fecha_anterior if fecha >= fecha_inicio)
    fechas_relevantes.extend(fecha for fecha in movimientos if fecha >= fecha_inicio)
    fecha_fin = max(fechas_relevantes)

    bases = [fecha for fecha in por_fecha_anterior if fecha < fecha_inicio]
    if bases:
        fecha_previa = max(bases)
        posicion_previa = _decimal(por_fecha_anterior[fecha_previa]["POSICION_FINAL_USD"])
        existe_previa = True
    else:
        fecha_previa = None
        posicion_previa = Decimal("0")
        existe_previa = False

    _desactivar_alertas_desde(conexion, posicion_id, fecha_inicio)
    deseadas: dict[date, dict] = {}
    fecha_actual = fecha_inicio
    while fecha_actual <= fecha_fin:
        if fecha_actual in anclas:
            ancla = anclas[fecha_actual]
            deseadas[fecha_actual] = ancla
            posicion_previa = _decimal(ancla["POSICION_FINAL_USD"])
            fecha_previa = fecha_actual
            existe_previa = True
            fecha_actual += timedelta(days=1)
            continue

        datos_mov = movimientos.get(fecha_actual)
        crear_inicial_cero = (
            not existe_previa
            and datos_mov is None
            and fecha_actual == fecha_minima_fin
        )
        if not existe_previa and datos_mov is None and not crear_inicial_cero:
            fecha_actual += timedelta(days=1)
            continue

        compras = datos_mov["COMPRAS"] if datos_mov else Decimal("0")
        ventas = datos_mov["VENTAS"] if datos_mov else Decimal("0")
        neto = compras - ventas
        final = posicion_previa + neto
        if crear_inicial_cero:
            origen = ORIGEN_INICIAL_CERO
            archivo = "INICIAL_CERO_PENDIENTE"
        else:
            origen = ORIGEN_MOVIMIENTOS if datos_mov else ORIGEN_CARRY
            archivo = datos_mov["ARCHIVO"] if datos_mov else "CARRY_FORWARD"
        ejecucion_fila = datos_mov["EJECUCION_ID"] if datos_mov else ejecucion_id

        if not existe_previa:
            if crear_inicial_cero:
                detalle_alerta = (
                    "La posicion activa no tiene saldo anterior ni movimientos en "
                    "el corte; se crea en cero para completar la tabla inicial. "
                    "El valor debe validarse o reemplazarse desde el Excel."
                )
            else:
                detalle_alerta = (
                    "No existe una posicion oficial anterior; se usa cero y el "
                    "proceso continua para mantener trazabilidad."
                )
            _insertar_alerta(
                conexion,
                ejecucion_id,
                _alerta(
                    "MISSING_PREVIOUS_POSITION",
                    fecha_actual.strftime("%d/%m/%Y"),
                    detalle_alerta,
                    posicion_id=posicion_id,
                ),
            )

        ahora = _ahora()
        deseadas[fecha_actual] = {
            "FECHA": fecha_actual.strftime("%d/%m/%Y"),
            "POSICION_ID": posicion_id,
            "POSICION_NOMBRE": definicion["POSICION_NOMBRE"],
            "FECHA_ANTERIOR": fecha_previa.strftime("%d/%m/%Y") if fecha_previa else None,
            "POSICION_ANTERIOR": _decimal_texto(posicion_previa),
            "COMPRAS_USD": _decimal_texto(compras),
            "VENTAS_USD": _decimal_texto(ventas),
            "MOVIMIENTO_NETO_USD": _decimal_texto(neto),
            "POSICION_FINAL_USD": _decimal_texto(final),
            "ARCHIVO_FUENTE": archivo,
            "EJECUCION_ID": ejecucion_fila,
            "TIPO_ORIGEN": origen,
            "ESTADO": "PENDIENTE_ALERTAS",
            "TIENE_ALERTA": 0,
            "DETALLE_ALERTA": "",
            "USUARIO": _usuario(),
            "CREATED_AT": ahora,
            "UPDATED_AT": ahora,
        }
        posicion_previa = final
        fecha_previa = fecha_actual
        existe_previa = True
        fecha_actual += timedelta(days=1)

    for fecha_fila, fila_anterior in por_fecha_anterior.items():
        if fecha_fila < fecha_inicio or fecha_fila in anclas:
            continue
        if fecha_fila not in deseadas:
            _insertar_auditoria(
                conexion,
                ejecucion_id,
                fila_anterior["FECHA"],
                posicion_id,
                fila_anterior["POSICION_FINAL_USD"],
                None,
                motivo,
            )
            conexion.execute(
                f"DELETE FROM {TABLA_HISTORICO_SPOT} WHERE FECHA = ? AND POSICION_ID = ?",
                (fila_anterior["FECHA"], posicion_id),
            )

    for fecha_fila, fila_nueva in deseadas.items():
        if fecha_fila in anclas:
            continue
        fila_anterior = por_fecha_anterior.get(fecha_fila)
        anterior = fila_anterior["POSICION_FINAL_USD"] if fila_anterior else None
        nuevo = fila_nueva["POSICION_FINAL_USD"]
        if fila_anterior is None or any(
            fila_anterior[campo] != fila_nueva[campo]
            for campo in (
                "POSICION_ANTERIOR",
                "COMPRAS_USD",
                "VENTAS_USD",
                "MOVIMIENTO_NETO_USD",
                "POSICION_FINAL_USD",
                "TIPO_ORIGEN",
            )
        ):
            _insertar_auditoria(
                conexion,
                ejecucion_id,
                fila_nueva["FECHA"],
                posicion_id,
                anterior,
                nuevo,
                motivo,
            )
        _upsert_historico(conexion, fila_nueva)


def _actualizar_estado_alertas(conexion, posiciones: set[str] | None = None) -> None:
    if posiciones:
        marcadores = ",".join("?" for _ in posiciones)
        filas = conexion.execute(
            f"SELECT FECHA, POSICION_ID FROM {TABLA_HISTORICO_SPOT} "
            f"WHERE POSICION_ID IN ({marcadores})",
            tuple(posiciones),
        ).fetchall()
    else:
        filas = conexion.execute(
            f"SELECT FECHA, POSICION_ID FROM {TABLA_HISTORICO_SPOT}"
        ).fetchall()
    for fecha_texto, posicion_id in filas:
        detalles = conexion.execute(
            f"""
            SELECT CODIGO, DETALLE FROM {TABLA_ALERTAS_SPOT}
            WHERE FECHA = ? AND POSICION_ID = ? AND VIGENTE = 1
            ORDER BY CODIGO, CREATED_AT
            """,
            (fecha_texto, posicion_id),
        ).fetchall()
        detalle = " | ".join(f"{codigo}: {texto}" for codigo, texto in detalles)
        conexion.execute(
            f"""
            UPDATE {TABLA_HISTORICO_SPOT}
            SET TIENE_ALERTA = ?, DETALLE_ALERTA = ?, ESTADO = ?
            WHERE FECHA = ? AND POSICION_ID = ?
            """,
            (
                1 if detalles else 0,
                detalle,
                "CON_ALERTA" if detalles else "OK",
                fecha_texto,
                posicion_id,
            ),
        )


def _resumen_desde_conexion(conexion, fecha_corte: date) -> pd.DataFrame:
    fecha_texto = fecha_corte.strftime("%d/%m/%Y")
    consulta = f"""
        SELECT FECHA, POSICION_ID, POSICION_NOMBRE, FECHA_ANTERIOR,
               POSICION_ANTERIOR, COMPRAS_USD, VENTAS_USD,
               MOVIMIENTO_NETO_USD, POSICION_FINAL_USD, TIPO_ORIGEN,
               ESTADO, TIENE_ALERTA, DETALLE_ALERTA, ARCHIVO_FUENTE,
               EJECUCION_ID, UPDATED_AT
        FROM {TABLA_HISTORICO_SPOT}
        WHERE FECHA = ?
        ORDER BY POSICION_NOMBRE
    """
    return pd.read_sql_query(consulta, conexion, params=(fecha_texto,))


def obtener_resumen_spot(
    fecha_trabajo: str | date,
    *,
    pruebas: bool = False,
    ruta_db: str | Path | None = None,
) -> pd.DataFrame:
    ruta = _resolver_ruta_db(ruta_db, pruebas=pruebas)
    with closing(conectar_sqlite(ruta)) as conexion:
        _crear_esquema(conexion)
        return _resumen_desde_conexion(conexion, _fecha(fecha_trabajo))


def arrastrar_spot_acumulado(
    fecha_trabajo: str | date,
    logger=None,
    *,
    pruebas: bool = False,
    ruta_db: str | Path | None = None,
) -> pd.DataFrame:
    """Publica el saldo Spot acumulado de un corte sin movimientos nuevos.

    Si el historico acumulado ya contiene todas las posiciones activas para la
    fecha, se limita a construir la tabla canonica. En caso contrario crea el
    carry-forward desde la ultima posicion anterior, con compras y ventas en
    cero y trazabilidad propia, sin requerir un Reporte de Caja.
    """
    fecha_corte = _fecha(fecha_trabajo)
    fecha_texto = fecha_corte.strftime("%d/%m/%Y")
    ruta_base = _resolver_ruta_db(ruta_db, pruebas=pruebas)
    por_book, por_posicion, hash_parametros = _cargar_parametros_spot()
    posiciones_activas = {
        posicion_id
        for posicion_id, definicion in por_posicion.items()
        if definicion["ACTIVO"]
    }

    with closing(conectar_sqlite(ruta_base)) as conexion:
        _crear_esquema(conexion)
        resumen = _resumen_desde_conexion(conexion, fecha_corte)
        presentes = set(resumen.get("POSICION_ID", pd.Series(dtype=str)).astype(str))

        if not posiciones_activas.issubset(presentes):
            faltantes_sin_base: list[str] = []
            for posicion_id in sorted(posiciones_activas):
                filas = _filas_historico_posicion(conexion, posicion_id)
                if not any(_fecha(fila["FECHA"]) < fecha_corte for fila in filas):
                    faltantes_sin_base.append(posicion_id)
            if faltantes_sin_base:
                raise ValueError(
                    "Spot no puede trasladarse al corte porque no existe una "
                    "posicion acumulada anterior para: "
                    + ", ".join(faltantes_sin_base)
                )

            ejecucion_id = (
                f"SPOT-CARRY-{fecha_corte:%Y%m%d}-{uuid.uuid4().hex[:12].upper()}"
            )
            inicio = _ahora()
            fuente = "CARRY_FORWARD"
            hash_fuente = hashlib.sha256(
                f"{fuente}|{fecha_texto}".encode("utf-8")
            ).hexdigest()
            conexion.execute("BEGIN IMMEDIATE")
            conexion.execute(
                f"""
                INSERT INTO {TABLA_EJECUCIONES_SPOT} (
                    EJECUCION_ID, FECHA_CORTE, ARCHIVO_FUENTE, RUTA_FUENTE,
                    SHA256_FUENTE, SHA256_PARAMETROS, TAMANO_BYTES,
                    MODIFICADO_FUENTE, ORIGEN, ESTADO, USUARIO, INICIO, FIN,
                    MENSAJE
                ) VALUES (?, ?, ?, '', ?, ?, 0, NULL, ?, 'EN_PROCESO', ?, ?, NULL, '')
                """,
                (
                    ejecucion_id,
                    fecha_texto,
                    fuente,
                    hash_fuente,
                    hash_parametros,
                    ORIGEN_CARRY,
                    _usuario(),
                    inicio,
                ),
            )

            for posicion_id in sorted(posiciones_activas):
                _recalcular_posicion(
                    conexion,
                    por_posicion[posicion_id],
                    fecha_corte,
                    fecha_corte,
                    ejecucion_id,
                    "CARRY_FORWARD_FECHA_NO_HABIL",
                )
            _actualizar_estado_alertas(conexion, posiciones_activas)
            resumen = _resumen_desde_conexion(conexion, fecha_corte)
            conexion.execute(
                f"""
                UPDATE {TABLA_EJECUCIONES_SPOT}
                SET ESTADO = 'COMPLETADA', FIN = ?, MENSAJE = ?
                WHERE EJECUCION_ID = ?
                """,
                (
                    _ahora(),
                    f"{len(resumen)} posiciones trasladadas sin movimientos",
                    ejecucion_id,
                ),
            )
            conexion.commit()
            registrar_log(
                logger,
                "Spot: se creo el carry-forward acumulado para "
                f"{fecha_texto} ({len(resumen)} posiciones).",
            )
        else:
            registrar_log(
                logger,
                "Spot: el acumulado ya contiene el corte "
                f"{fecha_texto} ({len(resumen)} posiciones); no se recalcula.",
            )

        canonica = _tabla_canonica(resumen, por_posicion)
        canonica.attrs["resumen_spot"] = resumen
        _guardar_tablas_auxiliares_spot(
            conexion,
            ruta_base,
            fecha_texto,
            canonica,
            resumen,
            por_book,
            logger=logger,
        )
        return canonica


def _tabla_canonica(
    resumen: pd.DataFrame,
    por_posicion: dict[str, dict],
) -> pd.DataFrame:
    if resumen.empty:
        return pd.DataFrame(columns=COLUMNAS_POSICION)
    filas: list[dict] = []
    for _, fila in resumen.iterrows():
        posicion_id = str(fila["POSICION_ID"])
        definicion = por_posicion.get(posicion_id)
        if not definicion or not definicion["ACTIVO"]:
            continue
        filas.append(
            {
                "FECHA": str(fila["FECHA"]),
                "PRODUCTO": "Spot",
                "BOOK": definicion["POSICION_NOMBRE"],
                # Se entrega texto decimal exacto. La capa transversal de
                # consolidacion aplica su contrato numerico legado al publicar.
                "POSICION": _decimal_texto(fila["POSICION_FINAL_USD"]),
                "MONEDA_POSICION": definicion["MONEDA_POSICION"],
                "LB_LT": definicion["LB_LT"],
                "INSTRUMENTO": definicion["INSTRUMENTO"],
                "COMPANY": "Colombia",
                "CLASIFICACION_CONTABLE": "NA",
                "BANKING_CVA_DVA": "Banking",
            }
        )
    return pd.DataFrame(filas, columns=COLUMNAS_POSICION)


def _guardar_tablas_auxiliares_spot(
    conexion,
    ruta_base: Path,
    fecha_texto: str,
    canonica: pd.DataFrame,
    resumen: pd.DataFrame,
    por_book: dict[str, dict],
    logger=None,
) -> None:
    """Replica los soportes Spot sin alterar las tablas transaccionales actuales."""
    consultas = {
        "compras_ventas_por_book": (
            f"SELECT * FROM {TABLA_COMPRAS_VENTAS_BOOK_SPOT} WHERE FECHA = ?",
            (fecha_texto,),
        ),
        "movimientos": (
            f"SELECT * FROM {TABLA_MOVIMIENTOS_SPOT} WHERE FECHA = ? AND VIGENTE = 1",
            (fecha_texto,),
        ),
        "alertas": (
            f"SELECT * FROM {TABLA_ALERTAS_SPOT} WHERE FECHA = ? AND VIGENTE = 1",
            (fecha_texto,),
        ),
        "ejecuciones": (
            f"SELECT * FROM {TABLA_EJECUCIONES_SPOT} WHERE FECHA_CORTE = ?",
            (fecha_texto,),
        ),
        "auditoria": (
            f"SELECT * FROM {TABLA_AUDITORIA_SPOT} WHERE FECHA = ?",
            (fecha_texto,),
        ),
    }
    tablas = {
        "posicion_canonica": canonica,
        "posicion_acumulada": resumen,
        "parametros_book": pd.DataFrame(list(por_book.values())),
    }
    tablas.update(
        {
            nombre: pd.read_sql_query(sql, conexion, params=params)
            for nombre, (sql, params) in consultas.items()
        }
    )

    # ADICION AUXILIAR: al usar una ruta de prueba temporal, su auxiliar queda
    # en esa misma carpeta; en operación normal resuelve a risko_auxiliar.db.
    guardar_tablas_auxiliares_modulo(
        "Spot",
        fecha_texto,
        tablas,
        ruta_db=ruta_base.with_name("risko_auxiliar.db"),
        logger=logger,
    )


def ejecutar_spot(
    fecha_trabajo: str,
    logger=None,
    *,
    pruebas: bool = False,
    ruta_db: str | Path | None = None,
    ruta_insumo: str | Path | None = None,
) -> pd.DataFrame:
    """Procesa el Reporte de Caja y retorna la tabla canonica de Spot.

    El resumen enriquecido que usa la interfaz queda en
    ``resultado.attrs['resumen_spot']`` para conservar el contrato historico de
    retorno (DataFrame canonico) con el resto de Position Monitor.
    """
    fecha_corte = _fecha(fecha_trabajo)
    fecha_texto = fecha_corte.strftime("%d/%m/%Y")
    ruta_fuente = Path(ruta_insumo) if ruta_insumo is not None else _ruta_insumo(fecha_corte)
    if not ruta_fuente.exists():
        raise FileNotFoundError(
            "Vector no dejo el Reporte de Caja para Spot: " f"{ruta_fuente}"
        )

    configuracion = _cargar_configuracion_spot()
    por_book, por_posicion, hash_parametros_libros = _cargar_parametros_spot()
    hash_parametros = hashlib.sha256(
        (
            hash_parametros_libros
            + "|"
            + "|".join(sorted(configuracion["owner_tables_incluidas"]))
            + "|"
            + VERSION_LOGICA_SPOT
        ).encode("utf-8")
    ).hexdigest()
    hash_fuente = _hash_archivo(ruta_fuente)
    ruta_base = _resolver_ruta_db(ruta_db, pruebas=pruebas)

    registrar_log(logger, f"Spot: fuente {ruta_fuente}")
    registrar_log(logger, f"Spot: base {'de pruebas' if pruebas else 'oficial'} {ruta_base}")

    # La validacion del archivo ocurre antes de crear tablas. Si el encabezado
    # o los datos son invalidos, la base queda completamente intacta.
    if ruta_base.exists():
        with closing(conectar_sqlite(ruta_base)) as conexion:
            existe_esquema = conexion.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
                (TABLA_EJECUCIONES_SPOT,),
            ).fetchone()
            repetida = None
            if existe_esquema:
                repetida = conexion.execute(
                    f"""
                    SELECT EJECUCION_ID FROM {TABLA_EJECUCIONES_SPOT}
                    WHERE FECHA_CORTE = ? AND SHA256_FUENTE = ?
                      AND SHA256_PARAMETROS = ? AND ESTADO = 'COMPLETADA'
                    ORDER BY FIN DESC LIMIT 1
                    """,
                    (fecha_texto, hash_fuente, hash_parametros),
                ).fetchone()
            if repetida:
                ejecucion_duplicada = (
                    f"SPOT-DUP-{fecha_corte:%Y%m%d}-{uuid.uuid4().hex[:12].upper()}"
                )
                instante = _ahora()
                modificado_fuente = datetime.fromtimestamp(
                    ruta_fuente.stat().st_mtime
                ).astimezone().isoformat(timespec="seconds")
                _crear_esquema(conexion)
                conexion.execute("BEGIN IMMEDIATE")
                _guardar_compras_ventas_por_book(
                    conexion,
                    fecha_texto,
                    por_book,
                    ruta_fuente.name,
                    repetida[0],
                )
                conexion.execute(
                    f"""
                    INSERT INTO {TABLA_EJECUCIONES_SPOT} (
                        EJECUCION_ID, FECHA_CORTE, ARCHIVO_FUENTE, RUTA_FUENTE,
                        SHA256_FUENTE, SHA256_PARAMETROS, TAMANO_BYTES,
                        MODIFICADO_FUENTE, ORIGEN, ESTADO, USUARIO, INICIO, FIN,
                        MENSAJE
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'IDEMPOTENCIA', 'OMITIDA',
                              ?, ?, ?, ?)
                    """,
                    (
                        ejecucion_duplicada,
                        fecha_texto,
                        ruta_fuente.name,
                        str(ruta_fuente),
                        hash_fuente,
                        hash_parametros,
                        ruta_fuente.stat().st_size,
                        modificado_fuente,
                        _usuario(),
                        instante,
                        instante,
                        f"Archivo ya procesado en {repetida[0]}",
                    ),
                )
                _insertar_alerta(
                    conexion,
                    ejecucion_duplicada,
                    _alerta(
                        "DUPLICATE_SOURCE_FILE",
                        fecha_texto,
                        f"Archivo y parametros ya procesados en {repetida[0]}; "
                        "no se duplicaron movimientos ni posiciones.",
                    ),
                )
                conexion.commit()
                registrar_log(
                    logger,
                    "ALERTA DUPLICATE_SOURCE_FILE: el archivo y los parametros ya "
                    f"fueron procesados en {repetida[0]}; no se duplican registros.",
                )
                resumen = _resumen_desde_conexion(conexion, fecha_corte)
                canonica = _tabla_canonica(resumen, por_posicion)
                canonica.attrs["resumen_spot"] = resumen
                # ADICION AUXILIAR: refresca la foto aunque el archivo sea idempotente.
                _guardar_tablas_auxiliares_spot(
                    conexion,
                    ruta_base,
                    fecha_texto,
                    canonica,
                    resumen,
                    por_book,
                    logger=logger,
                )
                return canonica

    movimientos, alertas = _leer_movimientos(
        ruta_fuente,
        fecha_corte,
        por_book,
        configuracion["owner_tables_incluidas"],
        logger=logger,
    )

    ejecucion_id = f"SPOT-{fecha_corte:%Y%m%d}-{uuid.uuid4().hex[:12].upper()}"
    inicio = _ahora()
    usuario = _usuario()
    modificado = datetime.fromtimestamp(ruta_fuente.stat().st_mtime).astimezone().isoformat(
        timespec="seconds"
    )

    with closing(conectar_sqlite(ruta_base)) as conexion:
        _crear_esquema(conexion)
        conexion.execute("BEGIN IMMEDIATE")
        anteriores = conexion.execute(
            f"""
            SELECT EJECUCION_ID, SHA256_FUENTE FROM {TABLA_EJECUCIONES_SPOT}
            WHERE FECHA_CORTE = ? AND ESTADO = 'COMPLETADA'
            ORDER BY FIN
            """,
            (fecha_texto,),
        ).fetchall()
        fechas_posteriores = conexion.execute(
            f"SELECT DISTINCT FECHA FROM {TABLA_HISTORICO_SPOT}"
        ).fetchall()
        es_historico = any(_fecha(fila[0]) > fecha_corte for fila in fechas_posteriores)

        conexion.execute(
            f"""
            INSERT INTO {TABLA_EJECUCIONES_SPOT} (
                EJECUCION_ID, FECHA_CORTE, ARCHIVO_FUENTE, RUTA_FUENTE,
                SHA256_FUENTE, SHA256_PARAMETROS, TAMANO_BYTES,
                MODIFICADO_FUENTE, ORIGEN, ESTADO, USUARIO, INICIO, FIN, MENSAJE
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'EN_PROCESO', ?, ?, NULL, '')
            """,
            (
                ejecucion_id,
                fecha_texto,
                ruta_fuente.name,
                str(ruta_fuente),
                hash_fuente,
                hash_parametros,
                ruta_fuente.stat().st_size,
                modificado,
                ORIGEN_MOVIMIENTOS,
                usuario,
                inicio,
            ),
        )

        # Las versiones anteriores se conservan para auditoria, pero dejan de
        # participar en el calculo oficial.
        conexion.execute(
            f"UPDATE {TABLA_MOVIMIENTOS_SPOT} SET VIGENTE = 0 "
            "WHERE FECHA = ? AND VIGENTE = 1",
            (fecha_texto,),
        )
        conexion.execute(
            f"UPDATE {TABLA_ALERTAS_SPOT} SET VIGENTE = 0 "
            "WHERE FECHA = ? AND VIGENTE = 1",
            (fecha_texto,),
        )

        if anteriores and any(hash_anterior != hash_fuente for _, hash_anterior in anteriores):
            alertas.append(
                _alerta(
                    "SOURCE_FILE_CHANGE",
                    fecha_texto,
                    "El corte ya tenia otra version del archivo fuente; se recalcula "
                    "esta fecha y todo el historico posterior afectado.",
                )
            )
        if es_historico:
            alertas.append(
                _alerta(
                    "HISTORICAL_REPROCESSING",
                    fecha_texto,
                    "Se reprocesa una fecha anterior al ultimo historico disponible.",
                )
            )

        creado = _ahora()
        conexion.executemany(
            f"""
            INSERT INTO {TABLA_MOVIMIENTOS_SPOT} (
                MOVIMIENTO_ID, EJECUCION_ID, FECHA, SOURCE_ROW_ID, TRADE_ID,
                PRODUCTO_ORIGEN, BOOK_ORIGINAL, POSICION_ID, POSICION_NOMBRE,
                DESK, EVTYPE, SETTLE_CCY, AMOUNT_ORIGINAL, TIPO_MOVIMIENTO,
                MONTO_USD, ARCHIVO_FUENTE, FILA_FUENTE, PUBLICABLE, VIGENTE,
                CREATED_AT
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
            """,
            [
                (
                    str(uuid.uuid4()),
                    ejecucion_id,
                    mov["FECHA"],
                    mov["SOURCE_ROW_ID"],
                    mov["TRADE_ID"],
                    mov["PRODUCTO_ORIGEN"],
                    mov["BOOK_ORIGINAL"],
                    mov["POSICION_ID"],
                    mov["POSICION_NOMBRE"],
                    mov["DESK"],
                    mov["EVTYPE"],
                    mov["SETTLE_CCY"],
                    mov["AMOUNT_ORIGINAL"],
                    mov["TIPO_MOVIMIENTO"],
                    mov["MONTO_USD"],
                    ruta_fuente.name,
                    mov["FILA_FUENTE"],
                    mov["PUBLICABLE"],
                    creado,
                )
                for mov in movimientos
            ],
        )
        _guardar_compras_ventas_por_book(
            conexion,
            fecha_texto,
            por_book,
            ruta_fuente.name,
            ejecucion_id,
        )
        for alerta in alertas:
            _insertar_alerta(conexion, ejecucion_id, alerta)

        posiciones_con_movimiento = {
            mov["POSICION_ID"]
            for mov in movimientos
            if mov["PUBLICABLE"] and mov["POSICION_ID"]
        }
        posiciones_con_historia = {
            fila[0]
            for fila in conexion.execute(
                f"SELECT DISTINCT POSICION_ID FROM {TABLA_HISTORICO_SPOT}"
            ).fetchall()
            if fila[0] in por_posicion and por_posicion[fila[0]]["ACTIVO"]
        }
        # La salida siempre incluye el universo completo de posiciones activas.
        # Si una posicion no tiene historia ni movimientos, _recalcular_posicion
        # crea una fila inicial en cero, con alerta persistente de validacion.
        posiciones_activas = {
            posicion_id
            for posicion_id, definicion in por_posicion.items()
            if definicion["ACTIVO"]
        }
        afectadas = posiciones_con_movimiento | posiciones_con_historia | posiciones_activas

        for posicion_id in sorted(afectadas):
            anteriores_pos = _filas_historico_posicion(conexion, posicion_id)
            fechas_previas = [
                _fecha(fila["FECHA"])
                for fila in anteriores_pos
                if _fecha(fila["FECHA"]) < fecha_corte
            ]
            if fechas_previas:
                ultima = max(fechas_previas)
                inicio_recalculo = min(fecha_corte, ultima + timedelta(days=1))
            else:
                inicio_recalculo = fecha_corte
            _recalcular_posicion(
                conexion,
                por_posicion[posicion_id],
                inicio_recalculo,
                fecha_corte,
                ejecucion_id,
                "REPROCESO_MOVIMIENTOS" if anteriores else "CARGA_MOVIMIENTOS",
            )

        _actualizar_estado_alertas(conexion, afectadas)
        resumen = _resumen_desde_conexion(conexion, fecha_corte)
        conexion.execute(
            f"""
            UPDATE {TABLA_EJECUCIONES_SPOT}
            SET ESTADO = 'COMPLETADA', FIN = ?, MENSAJE = ?
            WHERE EJECUCION_ID = ?
            """,
            (
                _ahora(),
                f"{len(movimientos)} movimientos; {len(resumen)} posiciones; "
                f"{len(alertas)} alertas de fuente/mapeo",
                ejecucion_id,
            ),
        )
        conexion.commit()

        # ADICION AUXILIAR: copia todas las tablas Spot a la base separada.
        # La transacción y el esquema Spot existentes ya quedaron confirmados.
        canonica = _tabla_canonica(resumen, por_posicion)
        canonica.attrs["resumen_spot"] = resumen
        _guardar_tablas_auxiliares_spot(
            conexion,
            ruta_base,
            fecha_texto,
            canonica,
            resumen,
            por_book,
            logger=logger,
        )

    registrar_log(
        logger,
        f"Spot completado: {len(resumen)} posiciones acumuladas; ejecucion {ejecucion_id}.",
    )
    return canonica


def registrar_posiciones_iniciales(
    fecha_trabajo: str | date,
    posiciones: dict[str, Decimal | str | int],
    *,
    pruebas: bool = False,
    ruta_db: str | Path | None = None,
    archivo_fuente: str = "param_posiciones_iniciales_spot.xlsx",
    logger=None,
) -> pd.DataFrame:
    """Registra/reemplaza anclas manuales y recalcula el historico posterior."""
    fecha_inicial = _fecha(fecha_trabajo)
    fecha_texto = fecha_inicial.strftime("%d/%m/%Y")
    por_book, por_posicion, hash_parametros = _cargar_parametros_spot()
    lookup = dict(por_book)
    lookup.update({clave: valor for clave, valor in por_posicion.items()})

    normalizadas: dict[str, tuple[dict, Decimal]] = {}
    for nombre, valor in posiciones.items():
        definicion = lookup.get(normalizar_clave_param(nombre))
        if not definicion:
            raise ValueError(
                f"La posicion inicial {nombre!r} no existe para Spot en param_libros.csv."
            )
        if not definicion["ACTIVO"]:
            raise ValueError(f"La posicion inicial {nombre!r} esta inactiva.")
        normalizadas[definicion["POSICION_ID"]] = (definicion, _decimal(valor))
    if not normalizadas:
        raise ValueError("No se recibieron posiciones iniciales para Spot.")

    contenido = json.dumps(
        {clave: _decimal_texto(valor) for clave, (_, valor) in normalizadas.items()},
        sort_keys=True,
    )
    hash_fuente = hashlib.sha256(contenido.encode("utf-8")).hexdigest()
    ejecucion_id = f"SPOT-INI-{fecha_inicial:%Y%m%d}-{uuid.uuid4().hex[:12].upper()}"
    ruta_base = _resolver_ruta_db(ruta_db, pruebas=pruebas)
    ahora = _ahora()

    with closing(conectar_sqlite(ruta_base)) as conexion:
        _crear_esquema(conexion)
        repetida = conexion.execute(
            f"""
            SELECT EJECUCION_ID FROM {TABLA_EJECUCIONES_SPOT}
            WHERE FECHA_CORTE = ? AND SHA256_FUENTE = ?
              AND SHA256_PARAMETROS = ? AND ORIGEN = ? AND ESTADO = 'COMPLETADA'
            LIMIT 1
            """,
            (fecha_texto, hash_fuente, hash_parametros, ORIGEN_EXCEL),
        ).fetchone()
        if repetida:
            registrar_log(logger, f"Posiciones iniciales ya importadas en {repetida[0]}.")
            return _resumen_desde_conexion(conexion, fecha_inicial)

        conexion.execute("BEGIN IMMEDIATE")
        conexion.execute(
            f"""
            INSERT INTO {TABLA_EJECUCIONES_SPOT} (
                EJECUCION_ID, FECHA_CORTE, ARCHIVO_FUENTE, RUTA_FUENTE,
                SHA256_FUENTE, SHA256_PARAMETROS, TAMANO_BYTES,
                MODIFICADO_FUENTE, ORIGEN, ESTADO, USUARIO, INICIO, FIN, MENSAJE
            ) VALUES (?, ?, ?, ?, ?, ?, 0, ?, ?, 'EN_PROCESO', ?, ?, NULL, '')
            """,
            (
                ejecucion_id,
                fecha_texto,
                archivo_fuente,
                archivo_fuente,
                hash_fuente,
                hash_parametros,
                ahora,
                ORIGEN_EXCEL,
                _usuario(),
                ahora,
            ),
        )

        for posicion_id, (definicion, valor) in normalizadas.items():
            existente = conexion.execute(
                f"""
                SELECT POSICION_FINAL_USD FROM {TABLA_HISTORICO_SPOT}
                WHERE FECHA = ? AND POSICION_ID = ?
                """,
                (fecha_texto, posicion_id),
            ).fetchone()
            valor_anterior = existente[0] if existente else None
            fila = {
                "FECHA": fecha_texto,
                "POSICION_ID": posicion_id,
                "POSICION_NOMBRE": definicion["POSICION_NOMBRE"],
                "FECHA_ANTERIOR": None,
                "POSICION_ANTERIOR": "0",
                "COMPRAS_USD": "0",
                "VENTAS_USD": "0",
                "MOVIMIENTO_NETO_USD": "0",
                "POSICION_FINAL_USD": _decimal_texto(valor),
                "ARCHIVO_FUENTE": archivo_fuente,
                "EJECUCION_ID": ejecucion_id,
                "TIPO_ORIGEN": ORIGEN_EXCEL,
                "ESTADO": "OK",
                "TIENE_ALERTA": 0,
                "DETALLE_ALERTA": "",
                "USUARIO": _usuario(),
                "CREATED_AT": ahora,
                "UPDATED_AT": ahora,
            }
            _upsert_historico(conexion, fila)
            if valor_anterior != fila["POSICION_FINAL_USD"]:
                _insertar_auditoria(
                    conexion,
                    ejecucion_id,
                    fecha_texto,
                    posicion_id,
                    valor_anterior,
                    fila["POSICION_FINAL_USD"],
                    "INITIAL_POSITION_CHANGE",
                )
                if valor_anterior is not None:
                    _insertar_alerta(
                        conexion,
                        ejecucion_id,
                        _alerta(
                            "INITIAL_POSITION_CHANGE",
                            fecha_texto,
                            "Cambio de posicion inicial; se recalculo el historico posterior.",
                            posicion_id=posicion_id,
                        ),
                    )

            fechas_existentes = [
                _fecha(fila_db[0])
                for fila_db in conexion.execute(
                    f"SELECT FECHA FROM {TABLA_HISTORICO_SPOT} WHERE POSICION_ID = ?",
                    (posicion_id,),
                ).fetchall()
            ]
            fecha_fin = max(fechas_existentes) if fechas_existentes else fecha_inicial
            if fecha_fin > fecha_inicial:
                _recalcular_posicion(
                    conexion,
                    definicion,
                    fecha_inicial + timedelta(days=1),
                    fecha_fin,
                    ejecucion_id,
                    "INITIAL_POSITION_CHANGE",
                )

        _actualizar_estado_alertas(conexion, set(normalizadas))
        conexion.execute(
            f"""
            UPDATE {TABLA_EJECUCIONES_SPOT}
            SET ESTADO = 'COMPLETADA', FIN = ?, MENSAJE = ?
            WHERE EJECUCION_ID = ?
            """,
            (
                _ahora(),
                f"{len(normalizadas)} posiciones iniciales registradas",
                ejecucion_id,
            ),
        )
        resumen = _resumen_desde_conexion(conexion, fecha_inicial)
        conexion.commit()
    registrar_log(logger, f"Posiciones iniciales Spot registradas: {len(normalizadas)}.")
    return resumen


def _fecha_excel(valor: object) -> date:
    if isinstance(valor, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(valor))
    return _fecha(valor)


def importar_posiciones_iniciales_excel(
    *,
    pruebas: bool = False,
    ruta_db: str | Path | None = None,
    ruta_excel: str | Path | None = None,
    logger=None,
) -> pd.DataFrame:
    """Importa el Excel parametrico propio de Risko, nunca ReporteCaja."""
    configuracion = _cargar_configuracion_spot()
    ruta = Path(ruta_excel or configuracion["archivo_excel_posiciones"])
    hoja = str(configuracion["hoja_excel_posiciones"])
    if not ruta.exists():
        raise FileNotFoundError(f"No existe el Excel de posiciones Spot: {ruta}")

    aliases_fecha = {"FECHA", "FECHA_CORTE"}
    aliases_posicion = {"POSICION", "BOOK", "POSICION_NOMBRE"}
    aliases_valor = {"POSICION_FINAL_USD", "POSICION_FINAL", "VALOR_USD"}
    aliases_estado = {"ESTADO_PARAMETRO", "ESTADO"}
    filas: list[list[object]] = []
    extension = ruta.suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Se requiere openpyxl para leer el parametro Spot.") from exc
        libro = load_workbook(ruta, read_only=True, data_only=True)
        try:
            if hoja not in libro.sheetnames:
                raise ValueError(f"El Excel parametrico Spot no contiene la hoja {hoja!r}.")
            pagina = libro[hoja]
            for numero, fila in enumerate(pagina.iter_rows(values_only=True)):
                if numero >= 5000:
                    break
                filas.append(list(fila))
        finally:
            libro.close()
    elif extension == ".xlsb":
        # Compatibilidad de lectura; la ruta oficial parametrica es XLSX.
        try:
            from pyxlsb import open_workbook
        except ImportError as exc:
            raise RuntimeError("Se requiere pyxlsb para leer archivos .xlsb.") from exc
        with open_workbook(str(ruta)) as libro:
            if hoja not in libro.sheets:
                raise ValueError(f"El Excel parametrico Spot no contiene la hoja {hoja!r}.")
            with libro.get_sheet(hoja) as pagina:
                for numero, fila in enumerate(pagina.rows()):
                    if numero >= 5000:
                        break
                    filas.append([celda.v for celda in fila])
    else:
        raise ValueError(
            f"Formato no soportado para el parametro Spot: {ruta.suffix}. Use .xlsx."
        )

    encabezado = None
    indices = None
    indice_estado_parametro = None
    for numero, fila in enumerate(filas):
        claves = [normalizar_clave_param(valor) for valor in fila]
        indice_fecha = next((i for i, c in enumerate(claves) if c in aliases_fecha), None)
        indice_pos = next((i for i, c in enumerate(claves) if c in aliases_posicion), None)
        indice_valor = next((i for i, c in enumerate(claves) if c in aliases_valor), None)
        indice_estado = next((i for i, c in enumerate(claves) if c in aliases_estado), None)
        if None not in (indice_fecha, indice_pos, indice_valor):
            encabezado = numero
            indices = (indice_fecha, indice_pos, indice_valor)
            indice_estado_parametro = indice_estado
            break
    if encabezado is None or indices is None:
        raise ValueError(
            f"En la hoja {hoja!r} no se encontro la tabla parametrica. Se exigen "
            "los encabezados FECHA, POSICION y POSICION_FINAL_USD."
        )
    if indice_estado_parametro is None:
        raise ValueError(
            f"La tabla parametrica de {hoja!r} debe incluir ESTADO_PARAMETRO."
        )

    por_book, por_posicion, _ = _cargar_parametros_spot()
    lookup = dict(por_book)
    lookup.update(por_posicion)
    posiciones_esperadas = {
        posicion_id for posicion_id, datos in por_posicion.items() if datos["ACTIVO"]
    }
    por_fecha: dict[date, dict[str, Decimal]] = {}
    ids_por_fecha: dict[date, set[str]] = {}
    errores: list[str] = []
    i_fecha, i_pos, i_valor = indices
    for numero_fila, fila in enumerate(filas[encabezado + 1 :], start=encabezado + 2):
        if max(indices) >= len(fila):
            continue
        if fila[i_fecha] in (None, "") and fila[i_pos] in (None, ""):
            continue
        try:
            fecha_fila = _fecha_excel(fila[i_fecha])
        except ValueError as exc:
            errores.append(f"fila {numero_fila}: {exc}")
            continue
        posicion = str(fila[i_pos]).strip()
        if not posicion:
            errores.append(f"fila {numero_fila}: POSICION vacia")
            continue
        definicion = lookup.get(normalizar_clave_param(posicion))
        if not definicion or not definicion["ACTIVO"]:
            errores.append(
                f"fila {numero_fila}: posicion desconocida o inactiva {posicion!r}"
            )
            continue
        estado_parametro = (
            str(fila[indice_estado_parametro]).strip().upper()
            if indice_estado_parametro < len(fila)
            and fila[indice_estado_parametro] not in (None, "")
            else ""
        )
        estado_valido = estado_parametro == "VALIDADO"
        if not estado_valido:
            errores.append(
                f"fila {numero_fila}: ESTADO_PARAMETRO debe ser VALIDADO para "
                f"{definicion['POSICION_NOMBRE']}"
            )
        posicion_id = definicion["POSICION_ID"]
        if posicion_id in ids_por_fecha.setdefault(fecha_fila, set()):
            errores.append(
                f"fila {numero_fila}: posicion duplicada para {fecha_fila:%d/%m/%Y}: "
                f"{definicion['POSICION_NOMBRE']}"
            )
            continue
        try:
            valor = _decimal(fila[i_valor])
        except (InvalidOperation, ValueError):
            errores.append(
                f"fila {numero_fila}: POSICION_FINAL_USD vacia o invalida para "
                f"{definicion['POSICION_NOMBRE']}"
            )
            continue
        if not estado_valido:
            continue
        ids_por_fecha[fecha_fila].add(posicion_id)
        por_fecha.setdefault(fecha_fila, {})[definicion["POSICION_NOMBRE"]] = valor

    for fecha_fila, ids_presentes in ids_por_fecha.items():
        faltantes = posiciones_esperadas - ids_presentes
        if faltantes:
            nombres = [por_posicion[posicion_id]["POSICION_NOMBRE"] for posicion_id in sorted(faltantes)]
            errores.append(
                f"{fecha_fila:%d/%m/%Y}: faltan posiciones activas: {', '.join(nombres)}"
            )
    if errores:
        raise ValueError(
            "El Excel parametrico Spot no supera la validacion:\n- "
            + "\n- ".join(errores)
        )
    if not por_fecha:
        raise ValueError("La tabla parametrica Spot no contiene posiciones para importar.")

    resultado = pd.DataFrame()
    for fecha_fila in sorted(por_fecha):
        resultado = registrar_posiciones_iniciales(
            fecha_fila,
            por_fecha[fecha_fila],
            pruebas=pruebas,
            ruta_db=ruta_db,
            archivo_fuente=str(ruta),
            logger=logger,
        )
    return resultado


if __name__ == "__main__":
    ejecutar_spot(input("Ingrese fecha de trabajo (dd-mm-aaaa): ").strip())
